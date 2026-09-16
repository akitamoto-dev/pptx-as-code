#!/usr/bin/env python3
# /// script
# dependencies = ["python-pptx>=1.0"]
# ///
"""生成した pptx の「編集可能性」を検証する。納品前のゲートに使う。

- スライドごとに図形の種別を集計する（text / picture / autoshape / table / line / group）
- スライド全面を覆う画像（1 枚画像貼り付けの疑い。編集できない）を検出して警告する
- メディアに SVG が含まれていないか確認する
- 編集可能なテキストが 0 件のスライドを警告する

警告があれば終了コード 1。OOXML の規格違反（負のサイズなど）はここでは見ないので、
別途、作業ディレクトリの tools/normalize.py --check を通す。

使い方:
    python3 inspect_pptx.py <pptx> [--json <validation.json>]

依存は python-pptx だけ。`uv run inspect_pptx.py ...` なら自動で入る。
"""
import argparse
import json
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def main():
    ap = argparse.ArgumentParser(description="pptx の編集可能性を検証する")
    ap.add_argument("pptx")
    ap.add_argument("--json", default=None, help="結果を書き出す JSON のパス")
    a = ap.parse_args()

    prs = Presentation(a.pptx)
    sw, sh = prs.slide_width, prs.slide_height
    report = {"slide_width_emu": sw, "slide_height_emu": sh, "slides": [], "warnings": []}

    for i, slide in enumerate(prs.slides, 1):
        c = {"text": 0, "picture": 0, "autoshape": 0, "table": 0, "line": 0, "group": 0, "other": 0}
        fulls = []
        for shp in slide.shapes:
            if shp.has_text_frame and shp.text_frame.text.strip():
                c["text"] += 1
            st = shp.shape_type
            if st == MSO_SHAPE_TYPE.PICTURE:
                c["picture"] += 1
                try:
                    if shp.width >= sw * 0.92 and shp.height >= sh * 0.92:
                        fulls.append(shp.name)
                except Exception:  # noqa: BLE001
                    pass
            elif getattr(shp, "has_table", False):
                c["table"] += 1
            elif st == MSO_SHAPE_TYPE.GROUP:
                c["group"] += 1
            elif st == MSO_SHAPE_TYPE.AUTO_SHAPE:
                c["autoshape"] += 1
            elif st == MSO_SHAPE_TYPE.LINE:
                c["line"] += 1
            else:
                c["other"] += 1
        report["slides"].append({"index": i, "counts": c, "full_bleed_pictures": fulls})
        if fulls:
            report["warnings"].append(
                f"slide {i}: スライド全面を覆う画像があります（1 枚画像貼り付けの疑い・編集不可）: {fulls}"
            )
        if c["text"] == 0:
            report["warnings"].append(f"slide {i}: 編集可能テキストが 0 件です")

    try:
        for part in prs.part.package.iter_parts():
            if str(part.partname).lower().endswith(".svg"):
                report["warnings"].append(f"SVG メディアが含まれています: {part.partname}")
    except Exception:  # noqa: BLE001
        pass

    out = json.dumps(report, ensure_ascii=False, indent=2)
    if a.json:
        with open(a.json, "w", encoding="utf-8") as f:
            f.write(out)
        print(f"wrote {a.json}")
    print(out)
    print(f"警告 {len(report['warnings'])} 件")
    sys.exit(1 if report["warnings"] else 0)


if __name__ == "__main__":
    main()
